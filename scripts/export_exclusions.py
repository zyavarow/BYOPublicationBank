"""Spreadsheet of every paper suggested for exclusion or review, with reasons. Free: no API calls.
  .venv\\Scripts\\python.exe scripts\\export_exclusions.py --org "DADA2 Papers"
Categories (mutually exclusive, in this order): duplicate copy, DADA2 software paper, scanned/needs OCR,
possible non-paper. Possible duplicates (not subtracted anywhere) are listed on their own tab.
"""
import argparse
import os
import re
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from pubbank.db import psycopg  # noqa: E402  (must come before any other psycopg import)
import pandas as pd  # noqa: E402
from dotenv import load_dotenv  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

load_dotenv(ROOT / ".env")

SQL = """
SELECT p.filename, p.page_count, p.char_count, p.parse_status, coalesce(p.excluded,false), p.exclusion_reason,
       coalesce(p.needs_ocr,false), p.duplicate_of IS NOT NULL, d.filename, p.duplicate_reason,
       p.possible_non_paper, p.non_paper_reason, p.title, p.extracted_title, p.authors, p.year, p.journal,
       p.doi, p.pmid, p.pub_types, p.possible_duplicate_of IS NOT NULL, q.filename, p.possible_duplicate_reason
FROM papers p
LEFT JOIN papers d ON d.id = p.duplicate_of
LEFT JOIN papers q ON q.id = p.possible_duplicate_of
WHERE p.org_id = %s
ORDER BY p.filename
"""


def friendly_nonpaper(reason):
    parts = []
    for part in (reason or "").split(";"):
        part = part.strip()
        if part.startswith("PubMed type:"):
            parts.append("PubMed lists this as: " + part.split(":", 1)[1].strip() + " (not original research)")
        elif part.startswith("filename suggests"):
            parts.append("File name suggests a supplement, poster or slides")
        elif part == "very little text":
            parts.append("Very little text (likely a figure, abstract-only page or cover page)")
        elif part.startswith("no paper-like sections"):
            parts.append("No recognizable paper structure and no PubMed/Crossref match")
        elif part:
            parts.append(part)
    return "; ".join(parts)


def first_author(a):
    return None if not a else a[0] + (" et al." if len(a) > 1 else "")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--org", required=True)
    args = ap.parse_args()
    cols = ["filename", "pages", "chars", "status", "excluded", "excl_reason", "needs_ocr", "is_dup", "dup_of", "dup_reason",
            "possible_nonpaper", "nonpaper_reason", "title", "ext_title", "authors", "year", "journal", "doi", "pmid",
            "pub_types", "has_possible_dup", "possible_dup_of", "possible_dup_reason"]
    with psycopg.connect(os.environ["DATABASE_URL"]) as conn:
        org_id = conn.execute("SELECT id FROM organizations WHERE slug=%s", (args.org,)).fetchone()[0]
        rows = [dict(zip(cols, r)) for r in conn.execute(SQL, (org_id,)).fetchall()]

    out_rows, possible_dups = [], []
    for r in rows:
        base = {
            "File": r["filename"], "Title": r["title"] or r["ext_title"], "First author": first_author(r["authors"]),
            "Year": r["year"], "Journal": r["journal"], "DOI": r["doi"], "PMID": r["pmid"], "Pages": r["pages"],
        }
        if r["is_dup"]:
            cat, action = "1. Duplicate copy", "Exclude (keep the original)"
            why = f"Same paper as '{r['dup_of']}' ({(r['dup_reason'] or '').split(' as ')[0]})"
        elif r["excluded"]:
            cat, action = "2. About the DADA2 software, not the disease", "Exclude"
            m = re.search(r"sequencing terms (\d+), disease terms (\d+)", r["excl_reason"] or "")
            why = ("Reads as a DNA-sequencing / microbiome paper that uses or tests the DADA2 software"
                   + (f" ({m.group(1)} sequencing terms vs. {m.group(2)} disease terms in the text)" if m else ""))
        elif r["needs_ocr"]:
            cat, action = "3. Scanned, no usable text", "Needs OCR or a better copy"
            why = f"Pages are images or the text layer is unreadable ({r['pages']} pages, {r['chars'] or 0:,} characters extracted)"
        elif r["status"] == "ok" and r["possible_nonpaper"]:
            cat, action = "4. Possible non-paper", "Review (some may be worth keeping)"
            why = friendly_nonpaper(r["nonpaper_reason"])
        else:
            if r["has_possible_dup"]:
                possible_dups.append({**base, "Possibly the same as": r["possible_dup_of"],
                                      "Why flagged": (r["possible_dup_reason"] or "").split(";")[0],
                                      "YOUR DECISION": None, "YOUR COMMENTS": None})
            continue
        out_rows.append({"Category": cat, **base, "Why suggested": why,
                         "PubMed article type": ", ".join(r["pub_types"] or []) or None,
                         "Suggested action": action, "YOUR DECISION": None, "YOUR COMMENTS": None})

    df = pd.DataFrame(out_rows).sort_values(["Category", "File"])
    dups_df = pd.DataFrame(possible_dups, columns=["File", "Title", "First author", "Year", "Journal", "DOI", "PMID", "Pages",
                                                  "Possibly the same as", "Why flagged", "YOUR DECISION", "YOUR COMMENTS"])
    counts = df.groupby("Category").size().reset_index(name="Files")
    summary = pd.concat([counts, pd.DataFrame([{"Category": "TOTAL", "Files": int(counts.Files.sum())}])])
    summary["What it means"] = [
        "Exact copies of a paper already in the library. Safe to drop; the original stays.",
        "Papers that cite or test the DADA2 sequencing software; not about the disease. Suggested for exclusion; please confirm.",
        "PDFs with no readable text. They stay out until we OCR them or get a better copy.",
        "Editorials, comments, errata, supplements, agency reports, etc. Flagged by automatic rules, so some may be worth keeping.",
        "",
    ]
    out = ROOT / "data" / "reports" / f"suggested_exclusions_{args.org.replace(' ', '_')}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    colors = {"1": "E8E8E8", "2": "FBE9E7", "3": "FFF4D6", "4": "E3EEF9"}
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, frame in (("Summary", summary), ("All suggested", df), ("Possible duplicates", dups_df)):
            frame.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2" if name != "Summary" else "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2B6E5E")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(frame.columns, 1):
                w = max([len(str(col))] + [len(str(v)) for v in frame[col].head(150)])
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(w + 2, 8), 70)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        ws = xw.sheets["All suggested"]
        for row in ws.iter_rows(min_row=2):
            fill = colors.get(str(row[0].value)[:1])
            if fill:
                row[0].fill = PatternFill("solid", fgColor=fill)
        for name, frame in (("All suggested", df), ("Possible duplicates", dups_df)):
            ws = xw.sheets[name]
            col = list(frame.columns).index("YOUR DECISION") + 1
            letter = ws.cell(1, col).column_letter
            dv = DataValidation(type="list", formula1='"Keep,Exclude,Unsure"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{len(frame) + 200}")
            for r in range(1, len(frame) + 2):
                for cc in (col, col + 1):
                    ws.cell(r, cc).fill = PatternFill("solid", fgColor="FFFBE6")
    print(f"Wrote {out}")
    print(summary.to_string(index=False))
    print(f"Possible duplicates (not subtracted): {len(dups_df)}")


if __name__ == "__main__":
    main()
