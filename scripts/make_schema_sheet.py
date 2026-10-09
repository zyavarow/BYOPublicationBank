"""Writes the editable draft extraction schema to docs/extraction_schema_DADA2.xlsx (one-off helper)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "extraction_schema_DADA2.xlsx"

HEADER = ["Field", "Include? (Yes / No / Maybe)", "Type", "Allowed values or format", "Rule for the AI",
          "Notes from your library (416 clean papers)", "YOUR EDITS / COMMENTS"]

ROWS = [
    ["Study type", "Yes", "Pick one",
     "Case report/series; Cohort/registry; Clinical trial or treatment study; Animal study; In vitro (cells/biochemistry); "
     "Computational/structural; Review; Guideline/consensus; Other/off-topic",
     "Classify by what THIS paper itself did. If it did several things, pick the primary one.",
     "The VCP list (Cell/Animal/Patient/Computational) has no Review or Guideline; ~84 papers mention reviews, ~145 guidelines/consensus.", ""],
    ["Gene / protein focus", "Yes", "Pick one",
     "ADA2 (CECR1); ADA1 (ADA); Both; Other/none",
     "Which gene or enzyme the paper is about. Do not assume ADA means ADA2.",
     "The library includes ADA1 papers (e.g. PEG-ADA enzyme replacement); this keeps the two from being confused.", ""],
    ["Mutations studied", "Yes", "Tag list",
     "GENE + variant, one-letter form, e.g. ADA2 p.R169Q (also accept c.-notation if no protein change is given)",
     "Only variants in THIS paper's own patients or experiments. Do not list variants cited as background. "
     "Normalize Arg169Gln / p.Arg169Gln / R169Q to p.R169Q.",
     "Most mentioned: R169Q (104 papers), G47R (99), G47A (46), Y453C (45), G47V (32), G358R (26), H112Q (17), P251L (16). "
     "Papers write the same variant 3 different ways.", ""],
    ["Model system", "Yes", "Pick several",
     "Human patients; Healthy donors; Primary human cells; Human cell line; Mouse; Zebrafish; Other animal; Yeast/fungus; "
     "Recombinant protein; None (review/guideline)",
     "What the paper's own experiments used. Reviews and guidelines: None.",
     "Mouse ~150 papers mention; zebrafish ~42; macrophages ~188 and monocytes ~173 (primary cells).", ""],
    ["Techniques", "Yes", "Tag list",
     "Suggested vocabulary: ADA2 enzyme activity assay; Sanger sequencing; Whole-exome sequencing; Gene panel/NGS; Western blot; "
     "Flow cytometry; qPCR; ELISA; RNA-seq; Single-cell RNA-seq; Immunohistochemistry; Mass spectrometry; "
     "Cytokine/interferon signature; Imaging (MRI/MRA); Biopsy/histology; CRISPR; Crystallography/structure. Other = free text.",
     "Methods the paper itself performed, not ones it only cites.",
     "ADA2 activity assays dominate (~214 papers), then histology (~180), imaging (~134), Sanger (~92), Western blot (~84).", ""],
    ["Sample source", "Yes", "Tag list",
     "Serum; Plasma; PBMC; Whole blood; Bone marrow; Skin biopsy; Fibroblasts; CSF; Urine; Other tissue; Other = free text",
     "Where the samples analysed in this paper came from.",
     "Serum ~246, plasma ~217, bone marrow ~198, tissue biopsy ~140, CSF ~87.", ""],
    ["Reagents (incl. catalog #)", "Yes", "List of entries",
     "Each entry: name; vendor; catalog number; target or use (e.g. anti-ADA2 antibody; Abcam; ab12345; Western blot)",
     "Only reagents actually used in this paper's methods. Leave catalog number blank if not reported; never guess.",
     "~278 papers mention an antibody but only ~40 give anything like a catalog number, so many entries will have no catalog number. "
     "Why this matters: antibody differences can explain contradictory results between papers.", ""],
    ["Publication year", "Yes", "Number", "e.g. 2024",
     "Not extracted by AI: taken from PubMed/Crossref metadata already in the database (no API cost).",
     "Already resolved for ~90% of papers in Phase 1.", ""],
    ["Corresponding author(s)", "Yes", "List of entries",
     "Each entry: name; institution. e.g. Jane Smith; National Institutes of Health",
     "The author(s) the paper labels as corresponding author / correspondence, usually on the first page. "
     "If none is labelled, leave empty (do not guess the last author). No email addresses.",
     "Added at your request: tells you which PI or lab published a given report. Often the last author, but not always, "
     "so this is read from the paper itself, with its supporting quote.", ""],
    ["All authors", "Yes", "List of names",
     "Full author list in the order printed, e.g. Smith J; Lee K; Patel R",
     "Taken from PubMed/Crossref metadata already in the database (no API cost). "
     "For the ~10% of papers with no metadata match, read from the paper's title page instead.",
     "Added at your request.", ""],
    ["Last author", "Yes", "Name",
     "Last name in the author list",
     "Derived from the author list (no AI, no cost). Not asked of Claude.",
     "Added at your request: the senior/last author is often a corresponding author, but the first author can be too "
     "(that is normal). Only a corresponding author who is neither first nor last is flagged for a look.", ""],
    ["Treatments / interventions", "Yes", "Tag list",
     "Anti-TNF; IVIG; Stem-cell transplant (HSCT); Corticosteroids; Gene therapy; Enzyme replacement; JAK inhibitor; "
     "Plasma infusion (FFP); Other = free text",
     "Treatments given to THIS paper's patients or tested in its experiments, not ones mentioned in the background.",
     "Mentions: anti-TNF ~190 papers, IVIG ~179, HSCT ~169, steroids ~161, gene therapy ~103, enzyme replacement ~54, JAK inhibitors ~31. "
     "Patients and clinicians will likely ask about treatment.", ""],
    ["Clinical features reported", "Yes", "Tag list",
     "Vasculitis/PAN; Stroke; Cytopenia/marrow failure; Immunodeficiency; Livedo; Fever; Lymphoproliferation; "
     "Hepatosplenomegaly/portal hypertension; Other = free text",
     "Features observed in THIS paper's patients.",
     "Mentions: vasculitis ~258, stroke ~257, immunodeficiency ~252, cytopenia ~228, fever ~197, livedo ~160.", ""],
    ["Number of patients", "No", "Number", "Integer; blank if not a patient study",
     "Patients described or analysed in this paper (cases in a series, cohort size).",
     "Dropped at your request (kept here in case you change your mind).", ""],
    ["(your own field?)", "Maybe", "", "", "", "Add a row if there is something else you want pulled from every paper.", ""],
]

INSTRUCTIONS = [
    "HOW TO EDIT THIS FILE",
    "",
    "1. Change 'Include?' to Yes, No or Maybe for each field. (Maybe = you are not sure yet; I will ask.)",
    "2. Edit the allowed values, rules, or type directly in the cells, or write comments in the last column.",
    "3. Add rows at the bottom for any extra field you want. Delete values you do not want.",
    "4. Save the file (same name and place) and tell me. I turn it into the validated schema for the test run.",
    "",
    "WHAT HAPPENS FOR EVERY FIELD",
    "- Each extracted value is stored together with its supporting quote and the section/page it came from.",
    "- 'Not reported' is always an allowed answer; the AI is told never to guess.",
    "- I check in code that every quote really appears word for word in the paper.",
    "- Papers that are off-topic (not about DADA2) get 'Other/off-topic' / 'None' rather than made-up values.",
    "",
    "NOTE ON THE COUNTS",
    "The numbers in column F are how many papers MENTION a term anywhere outside the reference list, not how many studied it.",
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Schema"
    ws.append(HEADER)
    for r in ROWS:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B6E5E")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [26, 14, 16, 60, 48, 56, 36]
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[0].font = Font(bold=True)
        if row[1].value == "Maybe":
            row[1].fill = PatternFill("solid", fgColor="FBF3EE")
        if row[1].value == "No":
            for c in row[:6]:
                c.font = Font(color="9A968C", strike=True if c.column == 1 else False)
        row[6].fill = PatternFill("solid", fgColor="FFFBE6")
    ws.freeze_panes = "B2"
    dv = DataValidation(type="list", formula1='"Yes,No,Maybe"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{len(ROWS) + 30}")

    wi = wb.create_sheet("How to edit")
    for line in INSTRUCTIONS:
        wi.append([line])
    wi.column_dimensions["A"].width = 120
    wi["A1"].font = Font(bold=True, size=13)
    for cell in ("A8", "A14"):
        wi[cell].font = Font(bold=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
