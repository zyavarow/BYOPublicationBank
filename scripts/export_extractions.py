"""Write the extraction review workbook (XLSX) from the database. Free: no API calls.
  .venv\\Scripts\\python.exe scripts\\export_extractions.py --org "DADA2 Papers" --run-label v2test3
Sheets: How to review, Extractions (one row per paper; every category cell says Found / Not found / Not applicable
with confidence and a reason, color-coded), Categories (the same, one row per paper and category, easy to filter),
Evidence (every extracted value with its quote), Needs attention, Field guide (how the AI reads each category).
(The earlier v1 format is produced by export_extractions_v1.py.)
"""
import argparse
import os
import re
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from pubbank.db import psycopg  # noqa: E402  (must come before any other psycopg import)
import pandas as pd  # noqa: E402
from dotenv import load_dotenv  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

load_dotenv(ROOT / ".env")
from pubbank.template import SCHEMA_VERSION, active_categories, load_template  # noqa: E402

FILLS = {"High": "DDEFE3", "Medium": "FFF1CC", "Low": "F8D7D3", "NF": "FBE9DC", "NA": "E8E8E8"}


def surname(name):
    parts = re.sub(r"[^\w\s-]", " ", name or "").split()
    return parts[-1].lower() if parts else ""


def fmt_entry(cid, e):
    def d(key="detail"):
        return f" ({e[key]})" if e.get(key) else ""
    if cid == "mutations":
        return f"{e['gene']} {e['variant']}"
    if cid == "reagents":
        return (e["name"] + (f" [{e['vendor']}]" if e.get("vendor") else "")
                + (f" cat# {e['catalog_number']}" if e.get("catalog_number") else "")
                + (f" - {e['target_or_use']}" if e.get("target_or_use") else ""))
    if cid == "corresponding_authors":
        return e["name"] + (f" ({e['institution']})" if e.get("institution") else "")
    main = next((v for k, v in e.items() if k not in ("evidence", "detail") and isinstance(v, str)), "")
    return main + d()


def cell_text(cid, cat):
    conf = cat["confidence_final"]
    if cat["status"] == "Found":
        return "; ".join(fmt_entry(cid, e) for e in cat["entries"]) + f"  [{conf}]"
    if cat["status"] == "Not found":
        return f"NOT FOUND [{conf}]: {cat['rationale']}"
    return f"NOT APPLICABLE: {cat['rationale']}"


def cell_fill(cat):
    if cat["status"] == "Not applicable":
        return FILLS["NA"]
    if cat["status"] == "Not found":
        return FILLS["NF"]
    return FILLS[cat["confidence_final"]]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--org", required=True)
    ap.add_argument("--run-label", default="default")
    ap.add_argument("--template", help="path to the extraction template (for the Field guide)")
    args = ap.parse_args()
    t = load_template(args.template)
    cats = active_categories(t)
    with psycopg.connect(os.environ["DATABASE_URL"]) as conn:
        org_id = conn.execute("SELECT id FROM organizations WHERE slug=%s", (args.org,)).fetchone()[0]
        rows = conn.execute(
            """SELECT e.id, p.filename, p.title, p.journal, p.doi, p.pmid, e.status, e.error, e.result, e.derived,
                      e.quote_stats, e.cost_usd, e.template_version
               FROM extractions e JOIN papers p ON p.id = e.paper_id
               WHERE e.org_id=%s AND e.schema_version=%s AND e.run_label=%s ORDER BY p.filename""",
            (org_id, SCHEMA_VERSION, args.run_label)).fetchall()
        # standardized terms (HGVS / HPO) written by normalize_extractions.py, if it has been run
        terms = {}
        if rows and conn.execute("SELECT to_regclass('entry_terms')").fetchone()[0]:
            for eid, cat_id, idx, system, nid, label, score, method, details in conn.execute(
                    "SELECT extraction_id, category, entry_index, norm_system, norm_id, norm_label, score, method, details "
                    "FROM entry_terms WHERE extraction_id = ANY(%s) ORDER BY rank", ([r[0] for r in rows],)):
                if system == "HGVS":
                    parts = " / ".join(x for x in ((details or {}).get("hgvs_p"), (details or {}).get("hgvs_c")) if x)
                    txt = (parts or "(not recognized as a variant)") + ("" if (details or {}).get("status") == "valid"
                                                                      else f" [{(details or {}).get('status')}]")
                else:
                    txt = f"{nid} {label}" + ("" if method == "exact-label" else f" ({method}, {score:.2f})") if nid else "(no HPO match)"
                terms.setdefault((eid, cat_id, idx), []).append(txt)
    if not rows:
        sys.exit(f"No {SCHEMA_VERSION} extractions found for run label '{args.run_label}'.")

    wide, long_rows, evidence, attention, fills = [], [], [], [], []
    for eid, fn, title, journal, doi, pmid, status, error, result, derived, qs, cost, tver in rows:
        row = {"File": fn, "Title": title, "Journal": journal, "DOI": doi, "PMID": pmid}
        if status != "ok" or not result:
            row["Extraction problem"] = f"{status}: {error}"
            wide.append(row)
            fills.append({})
            attention.append({"File": fn, "Issue": f"extraction {status}", "Detail": error})
            continue
        d = derived or {}
        row.update({"Year": d.get("publication_year"), "All authors": "; ".join(d.get("all_authors") or []),
                    "Author list source": d.get("author_source"), "Last author": d.get("last_author")})
        row_fills = {}
        for c in cats:
            cat = result[c["id"]]
            row[c["label"]] = cell_text(c["id"], cat)
            row_fills[c["label"]] = cell_fill(cat)
            long_rows.append({
                "File": fn, "Category": c["label"], "Status": cat["status"], "Confidence": cat["confidence_final"],
                "AI's own confidence": cat["confidence"], "Quote check": cat["evidence_check"],
                "How the AI read it (rationale)": cat["rationale"],
                "Result": "; ".join(fmt_entry(c["id"], e) for e in cat["entries"]),
                "Standardized (HGVS / HPO)": "; ".join("; ".join(terms.get((eid, c["id"], i), [])) for i in range(len(cat["entries"]))
                                                       if terms.get((eid, c["id"], i))),
                "Rule-based (not the AI's call)": "yes" if cat["rule_based"] else "",
                "Flags": " | ".join(cat["flags"])})
            if cat["flags"]:
                attention.append({"File": fn, "Issue": f"{c['label']}: AI answer was corrected", "Detail": " | ".join(cat["flags"])})
            if cat["status"] != "Not applicable" and cat["confidence_final"] == "Low":
                attention.append({"File": fn, "Issue": f"{c['label']}: low confidence ({cat['status']})",
                                  "Detail": cat["rationale"] or cat["evidence_check"]})
            for i, e in enumerate(cat["entries"]):
                ev = e["evidence"]
                evidence.append({"File": fn, "Category": c["label"], "Value": fmt_entry(c["id"], e),
                                 "Standardized term": "; ".join(terms.get((eid, c["id"], i), [])) or None,
                                 "Section": ev["section"], "Page": ev["page"] or None, "Quote": ev["quote"],
                                 "Quote check": ev.get("verification")})
                if ev.get("verification") in ("not_found", "too_short"):
                    attention.append({"File": fn, "Issue": f"quote {ev['verification']}: {c['label']}",
                                      "Detail": f"{fmt_entry(c['id'], e)} | \"{ev['quote'][:160]}\""})
        corr = result["corresponding_authors"]
        authors = d.get("all_authors") or []
        names = [e["name"] for e in corr["entries"]]
        if names and authors and not {surname(authors[0]), surname(authors[-1])} & {surname(n) for n in names}:
            attention.append({"File": fn, "Issue": "corresponding author is neither the first nor the last author",
                              "Detail": f"first: {authors[0]} | last: {authors[-1]} | corresponding: {', '.join(names)}"})
        row["Reviewer note from AI"] = result.get("extraction_notes")
        row["Quotes verified"] = f"{qs['verified']}/{qs['total']}" if qs else None
        row["Cost USD"] = float(cost) if cost is not None else None
        row["Template version"] = tver
        row["REVIEWER: correct? (Yes / Partly / No)"] = None
        row["REVIEWER: comments"] = None
        wide.append(row)
        fills.append(row_fills)

    guide = [{"Category": c["label"], "What the AI looks for": c["counts"], "What it ignores": c["not_counts"],
              "How it interprets it": c["interpret"],
              "Applies to study types": "All" if c["applies_to"] == "All" else "; ".join(c["applies_to"]),
              "'Not applicable' means": c["na_means"], "'Not found' means": c["nf_means"],
              "Usually found in": c["found_in"], "Confidence guide": c["confidence_guide"]} for c in cats]
    how_to = [
        "HOW TO REVIEW THIS WORKBOOK", "",
        "For each paper, an AI model read the text and filled in a fixed set of categories. Every category always has an answer: Found, Not found, or Not applicable.",
        "",
        "STATUS AND CONFIDENCE (colors on the Extractions tab)",
        "  Green / amber / red = Found with High / Medium / Low confidence.   Orange = NOT FOUND (the category applies but the paper does not state it).   Grey = NOT APPLICABLE (the category does not apply to this kind of paper, e.g. reagents for a review).",
        "  Confidence combines the AI's own rating with an automatic check: a value whose quote cannot be found word for word in the paper is capped at Low; a near match is capped at Medium.",
        "",
        "HOW TO READ IT",
        "1. 'Extractions' tab: one row per paper. Use the last two (yellow) columns to mark each row Yes / Partly / No and add comments.",
        "2. 'Categories' tab: the same information as a filterable list (one row per paper and category), including the AI's one-sentence reason for each answer.",
        "3. 'Evidence' tab: every extracted value with the exact quote, section and page it came from.",
        "4. 'Field guide' tab: how the AI is told to read each category: what counts, what is ignored, and what Not found / Not applicable mean. If a rule looks wrong, tell us and we change the template.",
        "5. 'Needs attention' tab: places worth a look (low confidence, corrected answers, quotes that could not be matched, unusual author situations).",
        "",
        "LIMITS TO KEEP IN MIND",
        "- Text was machine-extracted from PDFs, so tables and figures are only partly captured; values that appear only there may be missed (and will show as Not found).",
        "- Running the same paper twice can give small differences in judgment fields (for example how a technique is tagged). Core facts such as variants, catalog numbers and corresponding authors were stable in our checks.",
        "- Page numbers refer to the page where the quoted section starts and can be a page off.",
        "- Year, all authors and last author come from PubMed/Crossref, not from the AI.",
    ]

    out_dir = ROOT / "data" / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"extraction_review_v2_{args.org.replace(' ', '_')}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    sheets = {"Extractions": pd.DataFrame(wide), "Categories": pd.DataFrame(long_rows), "Evidence": pd.DataFrame(evidence),
              "Needs attention": pd.DataFrame(attention, columns=["File", "Issue", "Detail"]), "Field guide": pd.DataFrame(guide)}
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            for c in ws[1]:
                c.font = Font(bold=True)
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [min(len(str(v)), 80) for v in df[col].head(100)])
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 10), 60 if name != "Field guide" else 50)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        ws = xw.sheets["Extractions"]
        cols = list(sheets["Extractions"].columns)
        for r, rf in enumerate(fills, start=2):
            for label, color in rf.items():
                ws.cell(r, cols.index(label) + 1).fill = PatternFill("solid", fgColor=color)
        for i, col in enumerate(cols, 1):
            if str(col).startswith("REVIEWER"):
                for r in range(1, ws.max_row + 1):
                    ws.cell(r, i).fill = PatternFill("solid", fgColor="FFFBE6")
                ws.column_dimensions[ws.cell(1, i).column_letter].width = 26 if "correct" in col else 50
                if "correct" in col:
                    dv = DataValidation(type="list", formula1='"Yes,Partly,No"', allow_blank=True)
                    ws.add_data_validation(dv)
                    letter = ws.cell(1, i).column_letter
                    dv.add(f"{letter}2:{letter}{ws.max_row + 200}")
        ws = xw.sheets["Categories"]
        cl = list(sheets["Categories"].columns)
        for r in range(2, ws.max_row + 1):
            st, conf = ws.cell(r, cl.index("Status") + 1).value, ws.cell(r, cl.index("Confidence") + 1).value
            color = FILLS["NA"] if st == "Not applicable" else FILLS["NF"] if st == "Not found" else FILLS.get(conf)
            if color:
                ws.cell(r, cl.index("Status") + 1).fill = PatternFill("solid", fgColor=color)
                ws.cell(r, cl.index("Confidence") + 1).fill = PatternFill("solid", fgColor=color)
        wh = xw.book.create_sheet("How to review", 0)
        for line in how_to:
            wh.append([line])
        wh.column_dimensions["A"].width = 150
        wh["A1"].font = Font(bold=True, size=14)
        for r in (5, 9, 16):
            wh.cell(r, 1).font = Font(bold=True)
        for row in wh.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    total = sum(float(r[11] or 0) for r in rows)
    n_blank = sum(1 for r in long_rows if not r["Status"])
    print(f"Wrote {out}\n{len(rows)} papers, {len(long_rows)} category answers ({n_blank} without a status), "
          f"{len(evidence)} evidence rows, {len(attention)} items needing attention. Extraction cost in this run: ${total:.4f}")


if __name__ == "__main__":
    main()
