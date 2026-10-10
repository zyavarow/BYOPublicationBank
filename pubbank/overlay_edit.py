"""Safe editing of an organization's overlay workbook (orgs/<folder>/overlay.xlsx) for the Extraction Schema screen.

Every change is written to a temporary copy, checked by loading the full merged template (core + overlay) and building each
call group's response schema, and only then moved into place. The previous file is kept as overlay.xlsx.bak.
"""
import json
import os
import re
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pubbank.template import load_template

SCHEMA_CHAR_LIMIT = 9000        # the API rejects response schemas of roughly 9-10k characters; stay under
HEADERS = {
    "Context": ["Key", "Value"],
    "Category settings": ["Category", "Field", "Action (replace / append)", "Text"],
    "Entry field changes": ["Category", "Field", "Action (add / remove / replace / describe)", "Values (separate with ;)",
                            "Insert before (optional, for add)", "Kind (pick-list / text, optional)", "Description (for describe)"],
    "Custom call groups": ["Call group", "Label", "Sections sent to the AI (separate with ;)",
                           "If those sections hold fewer than this many characters, send the whole paper",
                           "Also send every section that starts on this page or earlier (0 = none)",
                           "Re-read the whole paper when a category is not found? (Always / Only if key sections are missing / Never)",
                           "Key sections (a re-read is skipped when the paper has any of these)"],
    "Custom categories": ["ID", "Label", "Include? (Yes / No)", "Answer type", "Applies to study types", "What counts", "What does NOT count",
                          "How to interpret it", "'Not applicable' means", "'Not found' means", "Usually found in", "Confidence guide",
                          "Show as a search filter? (Yes / No)", "Call group (1-3)"],
    "Custom entry fields": ["Category ID", "Field", "Kind (pick-list / text)", "Allowed values (separate with ;)", "Description"],
    "Rule additions": ["Rule", "Text appended to the global rules"],
    "Display groups": ["Category", "Group label shown on search chips", "Match (values, or ontology ids such as HP:0002633)"],
    "Genes": ["Symbol", "Aliases (separate with ;)", "Notes"],
}
CAT_KEYS = ["id", "label", "include", "answer_type", "applies_to", "counts", "not_counts", "interpret", "na_means", "nf_means", "found_in",
            "confidence_guide", "filter_chip", "call_group"]
RESERVED = {"extraction_notes", "all_authors_from_title_page"}


def new_overlay(path, disease):
    """Create an overlay with no differences from the standard template except the disease description."""
    wb = Workbook()
    wb.active.title = "Context"
    for name, header in HEADERS.items():
        ws = wb["Context"] if name == "Context" else wb.create_sheet(name)
        ws.append(header)
    wb["Context"].append(["disease", disease.strip()])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    load_template(path)          # fails loudly if the standard template cannot take this disease text


def slug_id(label, taken):
    base = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")[:40] or "field"
    if base[0].isdigit():
        base = "f_" + base
    cid, n = base, 2
    while cid in taken or cid in RESERVED:
        cid, n = f"{base}_{n}", n + 1
    return cid


def join(vals):
    return "; ".join(v for v in vals if v)


class Overlay:
    def __init__(self, path):
        self.path = Path(path)
        self.wb = load_workbook(self.path)
        for name, header in HEADERS.items():            # older overlays may lack some sheets
            if name not in self.wb.sheetnames:
                self.wb.create_sheet(name).append(header)

    # ---- generic row helpers
    def _rows(self, sheet):
        return list(self.wb[sheet].iter_rows(min_row=2))

    def _find(self, sheet, *keys):
        for row in self._rows(sheet):
            if all(str(row[i].value or "").strip().lower() == str(k).lower() for i, k in enumerate(keys)):
                return row
        return None

    def _delete(self, sheet, predicate):
        ws = self.wb[sheet]
        for row in reversed(self._rows(sheet)):
            if predicate(row):
                ws.delete_rows(row[0].row)

    def _upsert(self, sheet, keys, values):
        row = self._find(sheet, *keys)
        if row:
            for i, v in enumerate(values):
                row[i].value = v
        else:
            self.wb[sheet].append(list(values))

    # ---- edits
    def set_context(self, key, value):
        self._upsert("Context", (key,), (key, value))

    def custom_ids(self):
        return {str(r[0].value).strip() for r in self._rows("Custom categories") if r[0].value}

    def set_setting(self, cid, field, text, action="replace"):
        self._upsert("Category settings", (cid, field), (cid, field, action, text))

    def clear_setting(self, cid, field):
        self._delete("Category settings", lambda r: str(r[0].value) == cid and str(r[1].value) == field)

    def set_options(self, cid, field, values, kind=None):
        """Make a field's allowed values exactly `values` (replaces earlier add / remove / replace rows for that field)."""
        self._delete("Entry field changes", lambda r: str(r[0].value) == cid and str(r[1].value) == field
                     and str(r[2].value).lower() in ("add", "remove", "replace"))
        self.wb["Entry field changes"].append([cid, field, "replace", join(values), "", kind or "pick-list", ""])

    def set_custom_text(self, cid, key, value):
        """Change a text column of a custom category's own row."""
        col = CAT_KEYS.index(key)
        row = self._find("Custom categories", cid)
        if row is None:
            raise KeyError(cid)
        row[col].value = value

    def set_custom_options(self, cid, field, values):
        row = self._find("Custom entry fields", cid, field)
        if row is None:
            raise KeyError((cid, field))
        row[3].value = join(values)

    def remove_custom(self, cid):
        self._delete("Custom categories", lambda r: str(r[0].value) == cid)
        self._delete("Custom entry fields", lambda r: str(r[0].value) == cid)

    def add_custom(self, cat, fields):
        """cat: dict with CAT_KEYS (call_group may be '' to choose automatically). fields: list of dicts(name, kind, values, description)."""
        row = [cat.get(k, "") for k in CAT_KEYS]
        for f in fields:
            self.wb["Custom entry fields"].append([cat["id"], f["name"], f["kind"], join(f.get("values", [])), f.get("description", "")])
        gid = cat.get("call_group")
        if gid:
            self.wb["Custom categories"].append(row)
            return
        # choose a call group whose response schema stays small enough; create a new group if none has room
        ws = self.wb["Custom categories"]
        gi = CAT_KEYS.index("call_group")
        for gid in self._groups_in_order(cat.get("found_in", "")):
            row[gi] = gid
            ws.append(row)
            if self._fits(gid):
                return
            ws.delete_rows(ws.max_row)
        row[gi] = self._new_group()
        ws.append(row)
        if not self._fits(row[gi]):
            ws.delete_rows(ws.max_row)
            raise ValueError("Could not fit this field into any reading pass.")

    def _groups_in_order(self, found_in):
        t = load_template(self.path)
        ids = [g for g in t["groups"] if g != next(iter(t["groups"]))]       # never the first group (study type lives there)
        methods_first = "method" in found_in.lower()
        ids.sort(key=lambda g: ("methods" not in " ".join(t["groups"][g]["kinds"]), g) if methods_first else ("results" not in " ".join(t["groups"][g]["kinds"]), g))
        return ids

    def _new_group(self):
        t = load_template(self.path)
        gid = str(max(int(g) for g in t["groups"] if g.isdigit()) + 1)
        self.wb["Custom call groups"].append([gid, "Extra fields", "abstract; methods; results; case", 3000, 0, "Never", ""])
        return gid

    def _fits(self, gid):
        """Would the call group's response schema fit the API limit if the workbook were saved now?"""
        from pubbank.template import build_model
        tmp = self.path.with_suffix(".check.xlsx")
        self.wb.save(tmp)
        try:
            t = load_template(tmp)
            return len(json.dumps(build_model(t, gid).model_json_schema())) <= SCHEMA_CHAR_LIMIT
        finally:
            tmp.unlink(missing_ok=True)

    # ---- saving
    def save(self):
        """Validate the whole template, then replace the file (keeping a .bak). Raises ValueError with a plain message on failure."""
        tmp = self.path.with_suffix(".new.xlsx")
        self.wb.save(tmp)
        try:
            from pubbank.template import build_model
            t = load_template(tmp)
            for gid in t["groups"]:
                if any(c["include"] and c["call_group"] == gid for c in t["categories"]):
                    size = len(json.dumps(build_model(t, gid).model_json_schema()))
                    if size > SCHEMA_CHAR_LIMIT + 600:
                        raise ValueError(f"Reading pass {gid} would be too large for the AI ({size:,} characters). Remove a field or shorten its options.")
        except ValueError:
            tmp.unlink(missing_ok=True)
            raise
        except Exception as e:
            tmp.unlink(missing_ok=True)
            raise ValueError(f"That change makes the template invalid: {e}") from e
        shutil.copy2(self.path, self.path.with_suffix(".xlsx.bak"))
        os.replace(tmp, self.path)
