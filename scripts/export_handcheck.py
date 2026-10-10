"""Build the hand-check workbook: one row per paper and category, so a person can verify every AI answer.
Free (no API calls).
  .venv\\Scripts\\python.exe scripts\\export_handcheck.py --org "DADA2 Papers" --run-label v2final15
Sheets: Start here, Check each answer (verdict dropdown + notes), Missed items, Scorecard (live formulas),
Paper index (links to the PDFs), Definitions (how the AI was told to read each category).
"""
import argparse
import os
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pubbank.db import connect  # noqa: E402  (must come before any other psycopg import)
from dotenv import load_dotenv  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

load_dotenv(ROOT / ".env")
from export_extractions import FILLS, fmt_entry  # noqa: E402
from pubbank.template import SCHEMA_VERSION, active_categories, load_template  # noqa: E402

VERDICTS = ["Correct", "Correct but incomplete", "Partly wrong", "Wrong", "Unsure"]
YELLOW = PatternFill("solid", fgColor="FFFBE6")
HEAD = PatternFill("solid", fgColor="2B6E5E")


def bullets(items):
    return "\n".join(f"• {x}" for x in items if x)


def short_paper(authors, year, title):
    first = (authors[0].split()[-1] + (" et al." if len(authors) > 1 else "")) if authors else "?"
    return f"{first} {year or ''} – {(title or '')[:62]}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--org", required=True)
    ap.add_argument("--run-label", default="default")
    args = ap.parse_args()
    t = load_template()
    cats = active_categories(t)
    with connect() as conn:
        org_id = conn.execute("SELECT id FROM organizations WHERE slug=%s", (args.org,)).fetchone()[0]
        rows = conn.execute(
            """SELECT e.id, p.filename, p.relpath, p.title, p.journal, p.year, p.authors, p.doi, p.pmid, e.result
               FROM extractions e JOIN papers p ON p.id = e.paper_id
               WHERE e.org_id=%s AND e.schema_version=%s AND e.run_label=%s AND e.status='ok'
               ORDER BY p.filename""", (org_id, SCHEMA_VERSION, args.run_label)).fetchall()
        terms = {}
        if rows and conn.execute("SELECT to_regclass('entry_terms')").fetchone()[0]:
            for eid, cid, idx, system, nid, label, score, method, details in conn.execute(
                    "SELECT extraction_id, category, entry_index, norm_system, norm_id, norm_label, score, method, details "
                    "FROM entry_terms WHERE extraction_id = ANY(%s) ORDER BY rank", ([r[0] for r in rows],)):
                if system == "HGVS":
                    txt = " / ".join(x for x in ((details or {}).get("hgvs_p"), (details or {}).get("hgvs_c")) if x) or "(not recognized)"
                else:
                    txt = (f"{nid} {label}" + ("" if method == "exact-label" else f" ({method})")) if nid else "(no HPO match)"
                terms.setdefault((eid, cid, idx), []).append(txt)
    if not rows:
        sys.exit(f"No {SCHEMA_VERSION} extractions for run label '{args.run_label}'.")

    wb = Workbook()
    # ------------------------------------------------------------------ Check each answer
    ws = wb.active
    ws.title = "Check each answer"
    header = ["#", "Paper", "Open PDF", "Category", "AI status", "AI confidence", "What the AI answered",
              "Standardized term (HGVS / HPO)", "AI's reason", "Evidence (section, page, exact quote)",
              "YOUR VERDICT", "What is wrong or missing? (your note)", "Where in the paper? (page)"]
    ws.append(header)
    paper_labels, n = [], 0
    for eid, fn, relpath, title, journal, year, authors, doi, pmid, result in rows:
        label = short_paper(authors, year, title)
        paper_labels.append(label)
        pdf = ROOT / "data" / "raw" / args.org / relpath
        for c in cats:
            cat = result[c["id"]]
            n += 1
            entries = cat["entries"]
            page = next((e["evidence"]["page"] for e in entries if e["evidence"].get("page")), 1) or 1
            std = [t_ for i in range(len(entries)) for t_ in terms.get((eid, c["id"], i), [])]
            answer = bullets(fmt_entry(c["id"], e) for e in entries) if entries else \
                ("(nothing found)" if cat["status"] == "Not found" else "(does not apply to this kind of paper)")
            evidence = bullets(f"[{e['evidence']['section']}, p.{e['evidence']['page'] or '?'}] “{e['evidence']['quote']}”"
                               for e in entries) if entries else ""
            ws.append([n, label, "", c["label"], cat["status"], cat["confidence_final"], answer, bullets(std), cat["rationale"],
                       evidence, None, None, None])
            r = ws.max_row
            link = ws.cell(r, 3)
            link.value = f"Open (p.{page})"
            link.hyperlink = "file:///" + quote(str(pdf).replace("\\", "/"), safe="/:") + f"#page={page}"
            link.font = Font(color="0563C1", underline="single")
            color = FILLS["NA"] if cat["status"] == "Not applicable" else FILLS["NF"] if cat["status"] == "Not found" else FILLS[cat["confidence_final"]]
            ws.cell(r, 5).fill = ws.cell(r, 6).fill = PatternFill("solid", fgColor=color)
            for col in (11, 12, 13):
                ws.cell(r, col).fill = YELLOW
    last = ws.max_row
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD if cell.column < 11 else PatternFill("solid", fgColor="B5592E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate([5, 34, 11, 22, 13, 12, 52, 28, 42, 70, 22, 44, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:M{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{last}")

    # ------------------------------------------------------------------ Missed items
    wm = wb.create_sheet("Missed items")
    wm.append(["Paper", "Category", "What the paper states that the AI did not extract", "Where? (section / page)"])
    cat_labels = [c["label"] for c in cats]
    wl = wb.create_sheet("Lists")
    for i, p in enumerate(paper_labels, 1):
        wl.cell(i, 1, p)
    for i, cl in enumerate(cat_labels, 1):
        wl.cell(i, 2, cl)
    wl.sheet_state = "hidden"
    d1 = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(paper_labels)}", allow_blank=True)
    d2 = DataValidation(type="list", formula1=f"=Lists!$B$1:$B${len(cat_labels)}", allow_blank=True)
    wm.add_data_validation(d1); wm.add_data_validation(d2)
    d1.add("A2:A300"); d2.add("B2:B300")
    for cell in wm[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HEAD; cell.alignment = Alignment(wrap_text=True)
    for i, w in enumerate([44, 28, 90, 28], 1):
        wm.column_dimensions[get_column_letter(i)].width = w
    wm.freeze_panes = "A2"

    # ------------------------------------------------------------------ Scorecard (live formulas)
    sc = wb.create_sheet("Scorecard")
    rng = lambda col: f"'Check each answer'!${col}$2:${col}${last}"  # noqa: E731
    sc.append(["SCORECARD: fills in automatically as you choose verdicts on the 'Check each answer' tab"])
    sc["A1"].font = Font(bold=True, size=13)
    sc.append([])
    sc.append(["Overall", "Answers reviewed", "Correct", "Correct but incomplete", "Partly wrong", "Wrong", "Unsure", "% Correct"])

    def block(first_col_label, key_col, keys):
        for k in keys:
            r = sc.max_row + 1
            sc.append([k,
                       f'=COUNTIFS({rng(key_col)},A{r},{rng("K")},"<>")',
                       *[f'=COUNTIFS({rng(key_col)},$A{r},{rng("K")},"{v}")' for v in VERDICTS[:3]],
                       f'=COUNTIFS({rng(key_col)},$A{r},{rng("K")},"{VERDICTS[3]}")',
                       f'=COUNTIFS({rng(key_col)},$A{r},{rng("K")},"{VERDICTS[4]}")',
                       f'=IF(B{r}=0,"",C{r}/B{r})'])
            sc.cell(r, 8).number_format = "0%"

    r = sc.max_row + 1
    sc.append(["All answers", f'=COUNTA({rng("K")})', *[f'=COUNTIF({rng("K")},"{v}")' for v in VERDICTS],
               f'=IF(B{r}=0,"",C{r}/B{r})'])
    sc.cell(r, 8).number_format = "0%"
    sc.append([]); sc.append(["By category", "Answers reviewed", "Correct", "Correct but incomplete", "Partly wrong", "Wrong", "Unsure", "% Correct"])
    block("Category", "D", cat_labels)
    sc.append([]); sc.append(["By the AI's status", "Answers reviewed", "Correct", "Correct but incomplete", "Partly wrong", "Wrong", "Unsure", "% Correct"])
    block("Status", "E", ["Found", "Not found", "Not applicable"])
    sc.append([]); sc.append(["By the AI's confidence (does High really mean right?)", "Answers reviewed", "Correct", "Correct but incomplete", "Partly wrong", "Wrong", "Unsure", "% Correct"])
    block("Confidence", "F", ["High", "Medium", "Low"])
    for row in sc.iter_rows(min_row=3):
        if row[0].value in ("Overall", "By category", "By the AI's status") or str(row[0].value).startswith("By the AI"):
            for c in row:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD
    sc.column_dimensions["A"].width = 52
    for col in "BCDEFGH":
        sc.column_dimensions[col].width = 18

    # ------------------------------------------------------------------ Paper index
    wp = wb.create_sheet("Paper index")
    wp.append(["Paper", "Open PDF", "Journal", "Year", "DOI", "PMID", "File name"])
    for eid, fn, relpath, title, journal, year, authors, doi, pmid, result in rows:
        pdf = ROOT / "data" / "raw" / args.org / relpath
        wp.append([short_paper(authors, year, title), "Open PDF", journal, year, doi, pmid, fn])
        c = wp.cell(wp.max_row, 2)
        c.hyperlink = "file:///" + quote(str(pdf).replace("\\", "/"), safe="/:")
        c.font = Font(color="0563C1", underline="single")
    for cell in wp[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HEAD
    for i, w in enumerate([60, 12, 34, 7, 30, 11, 46], 1):
        wp.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------ Definitions
    wd = wb.create_sheet("Definitions")
    wd.append(["Category", "What the AI looks for", "What it ignores", "How it interprets it", "Applies to study types",
               "'Not applicable' means", "'Not found' means"])
    for c in cats:
        wd.append([c["label"], c["counts"], c["not_counts"], c["interpret"],
                   "All" if c["applies_to"] == "All" else "; ".join(c["applies_to"]), c["na_means"], c["nf_means"]])
    for cell in wd[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HEAD
    for i, w in enumerate([26, 46, 40, 70, 40, 36, 40], 1):
        wd.column_dimensions[get_column_letter(i)].width = w
    for row in wd.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # ------------------------------------------------------------------ Start here
    st = wb.create_sheet("Start here", 0)
    lines = [
        "HAND-CHECK GUIDE: DADA2 extraction, 15 papers",
        "",
        f"{len(rows)} papers x {len(cats)} categories = {n} answers to check. Plan on 2-3 hours; you can save and come back any time.",
        "",
        "WHAT TO DO",
        "1. Go to the 'Check each answer' tab. Each row is one AI answer for one category of one paper.",
        "2. Read 'What the AI answered', then use the quote in 'Evidence' to check it. Click 'Open (p.N)' to open the PDF (it may not jump to the page: use the page number shown; Ctrl+F the first words of the quote).",
        "3. Choose a verdict in the yellow 'YOUR VERDICT' column and, if it is not 'Correct', say what is wrong or missing in the next column.",
        "4. If the paper states something the AI did NOT extract, use the 'Missed items' tab (or write it in the note column).",
        "5. The 'Scorecard' tab fills in automatically: accuracy per category and, importantly, per confidence level.",
        "",
        "HOW TO JUDGE EACH KIND OF ANSWER",
        "- Found: 'Correct' = the values are right and nothing important is missing. 'Correct but incomplete' = right, but the paper has more. 'Partly wrong' / 'Wrong' = some / all of it is wrong.",
        "- Not found: 'Correct' = the paper really does not say it. 'Wrong' = the paper DOES say it (put the page in the last column).",
        "- Not applicable: 'Correct' = you agree the category does not apply to this kind of paper (for example reagents in a review). 'Wrong' = it should apply.",
        "- 'Unsure' = you cannot tell; leave a note.",
        "",
        "COLORS (AI status and confidence): green / amber / red = Found with High / Medium / Low confidence; orange = Not found; grey = Not applicable.",
        "",
        "THINGS WORTH KNOWING",
        "- 'Standardized term' shows how a variant was converted to the HGVS standard (p.Arg169Gln) or which Human Phenotype Ontology term a feature was mapped to. Judge the mapping too: is it the right term?",
        "- Quotes were checked by computer to appear word for word in the paper text, so a quote is real; the question is whether it supports the answer.",
        "- The 'Definitions' tab shows how the AI was told to read each category. If you disagree with a rule, note it: that is a template change, not just a one-off error.",
        "- Year, authors and journal come from PubMed, not from the AI, and are not part of this check.",
        "",
        "WHEN YOU ARE DONE: save the file and tell Claude. It can compute accuracy by category and by confidence from your verdicts, and list the patterns behind the errors.",
    ]
    for line in lines:
        st.append([line])
    st.column_dimensions["A"].width = 160
    st["A1"].font = Font(bold=True, size=14)
    for r in (5, 12, 18, 21):
        st.cell(r, 1).font = Font(bold=True)
    for row in st.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    out = ROOT / "data" / "reports" / f"handcheck_{args.org.replace(' ', '_')}_{args.run_label}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.move_sheet("Check each answer", offset=-(wb.sheetnames.index("Check each answer") - 1))
    wb.save(out)
    print(f"Wrote {out}\n{len(rows)} papers, {n} answers to check, {len(cats)} categories.")


if __name__ == "__main__":
    main()
