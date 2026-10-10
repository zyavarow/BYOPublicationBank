"""Seeds the first DRAFT of the VCP organization overlay (orgs/vcp/overlay.xlsx) from the onboarding interview
(orgs/vcp/interview.md) and a read of the 14 VCP sample papers. Free (no API calls). After seeding, the Excel file is the
source of truth: the project owner edits it. Refuses to overwrite unless --force.
  .venv\\Scripts\\python.exe scripts\\make_vcp_overlay.py
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_core_and_overlay import HEADER_CHANGES, HEADER_SETTINGS, style_sheet, write_how_to  # noqa: E402

DISEASE = ("VCP-related multisystem proteinopathy (MSP1; also called inclusion body myopathy with Paget disease of bone and "
           "frontotemporal dementia, IBMPFD), caused by variants in the VCP gene (the protein is VCP, also called p97; "
           "Cdc48 in yeast, TER94 in Drosophila). The library also covers VCP-linked ALS, parkinsonism and neuropathy, and "
           "the basic biology of VCP/p97 (structure, cofactors, autophagy, proteostasis, DNA repair). Closely related "
           "conditions such as sporadic inclusion body myositis, and other AAA+ ATPases, are NOT the subject")

SETTINGS = [
    ("study_type", "interpret", "replace",
     "If a paper does several things, pick the primary one (the one the title and abstract emphasize). Reviews, guidelines and "
     "consensus statements summarize other work. Use 'Other/off-topic' when the paper is not about VCP/p97 or the conditions "
     "named in the Context rule (for example sporadic inclusion body myositis, another AAA+ ATPase, or a paper that only "
     "matches the name); say so in the rationale. Structure and mechanism papers on purified or recombinant p97 (cryo-EM, "
     "biochemistry) are 'Computational/structural' if the paper's work is a structure or simulation, and 'In vitro "
     "(cells/biochemistry)' if it is activity or binding assays."),
    ("gene_focus", "counts", "replace",
     "Which gene or protein the paper is about: VCP (p97), a protein that works with VCP, another gene that causes a related "
     "multisystem proteinopathy, or something else."),
    ("gene_focus", "interpret", "replace",
     "VCP, p97 and Cdc48/TER94 (yeast/fly versions) are the same protein family for this bank. Choose 'VCP (p97)' when VCP is the "
     "main subject, including VCP mutant, inhibitor and structure work. Choose 'VCP cofactor or partner' when the paper is mainly "
     "about a protein that binds or regulates VCP (for example UFD1, NPL4, p47, UBXD1, FAF2) and VCP is the context. Choose "
     "'Other MSP gene' for HNRNPA2B1, HNRNPA1, SQSTM1 (p62), MATR3 or similar genes in related multisystem proteinopathies. "
     "Use 'More than one of these' only when the paper substantially studies both VCP and another listed protein. Papers on "
     "inclusion body myositis without a VCP link, or on unrelated genes, are 'Other/none'."),
    ("gene_focus", "confidence_guide", "replace",
     "High when the title or abstract names the protein; Medium when only inferable from context."),
    ("mutations", "interpret", "append",
     "VCP variants are written many ways (R155H, p.Arg155His, c.464G>A). Keep the paper's numbering exactly as written and do "
     "not renumber; if a paper says its numbering differs from the standard reference, say so in the rationale. Also record "
     "variants engineered into models (knock-in mice, Drosophila, iPSC lines)."),
    ("model_systems", "interpret", "append",
     "Put the organism or cell type in the pick-list. Specific strains, lines and genotypes go in the dedicated 'Model strains "
     "and cell lines' category; a short mention in the detail field is enough here."),
    ("treatments", "counts", "replace",
     "Treatments or management given to this paper's own patients (for example exercise, bisphosphonates for Paget disease, "
     "respiratory or cardiac support, trial drugs). Compounds tested in cells or animal models go in the 'Inhibitors and compounds "
     "tested' category instead."),
    ("treatments", "not_counts", "replace",
     "Treatments only mentioned in the background, in a review of the literature, or recommended by cited guidelines; compounds "
     "tested only in cells, flies, mice or purified protein (those belong to 'Inhibitors and compounds tested')."),
    ("clinical_features", "counts", "replace",
     "Clinical features observed in this paper's own patients, for example muscle weakness, myopathy, Paget disease of bone, "
     "frontotemporal dementia, ALS, parkinsonism, neuropathy, cardiomyopathy, respiratory failure. Features of animal models are "
     "not clinical features."),
]

FIELD_CHANGES = [
    ("gene_focus", "value", "replace",
     "VCP (p97); VCP cofactor or partner; Other MSP gene; More than one of these; Other/none", "", "pick-list", ""),
    ("gene_focus", "value", "describe", "", "", "", "The protein or gene the paper is mainly about. Support it with a sentence from the paper's own aims, methods or results, not a definition of the gene or disease."),
    ("mutations", "gene", "describe", "", "", "", "Gene symbol, for example VCP."),
    ("reagents", "name", "describe", "", "", "",
     "The reagent, for example 'anti-VCP antibody' or a specific clone."),
    ("techniques", "technique", "add", "ATPase activity assay; Cryo-EM", "Crystallography/structure", "", ""),
    ("model_systems", "system", "add", "Drosophila", "Other animal", "", ""),
    ("model_systems", "system", "add", "iPSC-derived cells", "Human cell line", "", ""),
    ("treatments", "treatment", "replace",
     "Small-molecule drug; Biologic/antibody; Gene therapy; Cell or stem-cell transplant; Exercise or physical therapy; "
     "Surgery/procedure; Supportive care; Other", "", "pick-list", ""),
]

CAT_HEADER = ["ID", "Label", "Include? (Yes / No)", "Answer type", "Applies to study types", "What counts", "What does NOT count",
              "How to interpret it", "'Not applicable' means", "'Not found' means", "Usually found in", "Confidence guide",
              "Show as a search filter? (Yes / No)", "Call group (1-3)"]
FIELD_HEADER = ["Category ID", "Field", "Kind (pick-list / text)", "Allowed values (separate with ;)", "Description"]

EXPERIMENTAL = ("Animal study; In vitro (cells/biochemistry); Computational/structural")
OWN_WORK = "Case report/series; Cohort/registry; Clinical trial or treatment study; " + EXPERIMENTAL

CUSTOM_CATS = [
    ["outcome_measures", "Outcome measures and biomarkers", "Yes", "Entries",
     "Case report/series; Cohort/registry; Clinical trial or treatment study; Animal study",
     "Tests, scales, imaging and laboratory markers this paper itself used to measure disease or treatment effect, with the "
     "values reported: strength tests (MRC, dynamometry), functional scales (6-minute walk, ALSFRS-R, FIM), pulmonary function, "
     "imaging (muscle MRI), cognitive tests, quality-of-life instruments, and blood or CSF biomarkers (creatine kinase, alkaline "
     "phosphatase, neurofilament light).",
     "Measures only cited from other studies or recommended in the background; routine diagnostic tests with no result reported.",
     "One entry per measure. Put the reported result (mean, range, change, time point) in 'result'; empty string if the paper "
     "names the measure but gives no value. Use the paper's own name for the instrument.",
     "Reviews, guidelines, structural or purely biochemical papers, and off-topic papers.",
     "A clinical or animal paper that reports no outcome measure or biomarker.",
     "Methods; Results; Tables", "High when the measure and its value are both stated; Medium when only the measure is named.",
     "Yes", "3"],
    ["compounds_tested", "Inhibitors and compounds tested", "Yes", "Entries",
     OWN_WORK,
     "Compounds this paper itself tested on VCP/p97 or in its models: VCP inhibitors (CB-5083, NMS-873, DBeQ, ML240, "
     "UPCDC30245), other drugs, and small molecules, with the target, dose and the system they were tested in.",
     "Compounds only discussed in the background, listed in a review, or recommended by guidelines; reagents used only as "
     "experimental tools (for example a protease inhibitor in a lysis buffer).",
     "One entry per compound. 'target' is what it acts on (for example VCP D2 ATPase domain); 'dose' as written with units; "
     "'tested_in' says the system (patients, mouse, Drosophila, iPSC-derived cells, purified protein). Put the main finding in "
     "'effect' in a few words.",
     "Reviews, guidelines and off-topic papers.",
     "A paper with its own experiments or trial that tests no compound.",
     "Abstract; Methods; Results", "High when compound and system are stated; Medium when the dose or target is missing.",
     "Yes", "4"],
    ["model_strains", "Model strains and cell lines", "Yes", "Entries",
     EXPERIMENTAL + "; Case report/series; Cohort/registry",
     "Specific mouse, fly, yeast or zebrafish lines and strains, and cell lines or iPSC lines, that this paper itself used, "
     "including their genotype (for example VCP R155H knock-in mouse, VCP R152H patient iPSC line, HEK293, C2C12).",
     "Models only cited from other studies; generic terms with no name or genotype.",
     "One entry per line or strain. 'genotype' is the VCP change or modification (as the paper writes it); 'source' is where it "
     "came from (company, repository, lab or patient-derived) and the identifier if printed (for example a Jackson Laboratory "
     "stock number or an RRID); empty strings when not reported.",
     "Reviews, guidelines, and papers with no models or cell lines (for example purified protein only, or a clinical paper).",
     "A paper using models or cell lines that does not name or describe them.",
     "Methods", "High when named and described in Methods; Medium when only mentioned in Results.",
     "Yes", "4"],
]

CUSTOM_FIELDS = [
    ("outcome_measures", "measure", "text", "", "The test, scale or biomarker, as the paper names it."),
    ("outcome_measures", "result", "text", "", "The value or change reported, with units and time point; empty string if none."),
    ("compounds_tested", "compound", "text", "", "The compound name as written (for example CB-5083)."),
    ("compounds_tested", "target", "text", "", "What it acts on; empty string if not stated."),
    ("compounds_tested", "dose", "text", "", "Dose or concentration as written, with units; empty string if not stated."),
    ("compounds_tested", "tested_in", "text", "", "The system it was tested in (patients, mouse, Drosophila, iPSC-derived cells, purified protein)."),
    ("compounds_tested", "effect", "text", "", "The main finding in a few words; empty string if none."),
    ("model_strains", "name", "text", "", "The strain or cell line name as the paper writes it."),
    ("model_strains", "organism_or_cell", "text", "", "Organism or cell type (mouse, Drosophila, human iPSC-derived motor neuron, and so on)."),
    ("model_strains", "genotype", "text", "", "The VCP change or other modification, as written; empty string for wild type or unstated."),
    ("model_strains", "source", "text", "", "Where it came from and any identifier printed (stock number, RRID); empty string if not reported."),
]

RULES = [
    ("Gene names", "For this bank VCP, p97 and valosin-containing protein are the same protein, and Cdc48 (yeast) and TER94 "
                   "(Drosophila) are its orthologs. Treat them as one when deciding what a paper is about, but record the name the "
                   "paper uses."),
    ("Sporadic IBM", "Sporadic inclusion body myositis (sIBM) is a different disease; a paper about it is off-topic unless it "
                     "studies VCP."),
]

DISPLAY_GROUPS = [
    ("clinical_features", "Muscle (myopathy, weakness)", "Myopathy; Muscle weakness; Proximal muscle weakness; Limb-girdle muscle weakness; Muscle atrophy; Inclusion body myositis"),
    ("clinical_features", "Bone (Paget disease)", "Paget disease of bone; Osteolysis; Abnormal bone structure"),
    ("clinical_features", "Dementia / cognition", "Frontotemporal dementia; Dementia; Cognitive impairment; Behavioral abnormality"),
    ("clinical_features", "ALS / motor neuron disease", "Amyotrophic lateral sclerosis; Motor neuron disease; Upper motor neuron dysfunction; Lower motor neuron dysfunction"),
    ("clinical_features", "Parkinsonism", "Parkinsonism"),
    ("clinical_features", "Neuropathy", "Peripheral neuropathy; Axonal neuropathy; Charcot-Marie-Tooth disease"),
    ("clinical_features", "Heart", "Cardiomyopathy; Dilated cardiomyopathy; Heart failure; Arrhythmia"),
    ("clinical_features", "Breathing", "Respiratory insufficiency; Respiratory failure; Dyspnea"),
]

CALL_GROUP_HEADER = [
    "Call group", "Label", "Sections sent to the AI (separate with ;)",
    "If those sections hold fewer than this many characters, send the whole paper",
    "Also send every section that starts on this page or earlier (0 = none)",
    "Re-read the whole paper when a category is not found? (Always / Only if key sections are missing / Never)",
    "Key sections (a re-read is skipped when the paper has any of these)"]
# The API rejects a response schema above roughly 9-10k characters, so the two extra categories that do not fit in groups
# 2 and 3 get their own reading pass. Re-reading the whole paper is switched off for this group to save cost.
CUSTOM_GROUPS = [
    ["4", "Models and compounds: model strains and cell lines, inhibitors and compounds tested",
     "abstract; methods; results; case", 3000, 0, "Never", ""],
]

GENES = [
    ("VCP", "p97; Cdc48; TER94; valosin-containing protein", "The gene mutated in VCP-related multisystem proteinopathy (MSP1). Cdc48 is the yeast ortholog and TER94 the Drosophila ortholog."),
    ("UFD1", "UFD1L", "VCP cofactor (with NPL4)."),
    ("NPLOC4", "NPL4", "VCP cofactor (with UFD1)."),
    ("NSFL1C", "p47", "VCP cofactor."),
    ("UBXN6", "UBXD1", "VCP cofactor."),
    ("FAF2", "UBXD8; ETEA", "VCP cofactor."),
    ("SQSTM1", "p62", "Gene for a related disease (Paget disease of bone, ALS); not VCP."),
    ("HNRNPA2B1", "hnRNPA2B1", "Gene for a related multisystem proteinopathy (MSP2); not VCP."),
    ("HNRNPA1", "hnRNPA1", "Gene for a related multisystem proteinopathy (MSP3); not VCP."),
    ("MATR3", "Matrin 3", "Gene for a related multisystem proteinopathy (MSP4); not VCP."),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    out = ROOT / "orgs" / "vcp" / "overlay.xlsx"
    if out.exists() and not args.force:
        sys.exit(f"{out} already exists; it is the source of truth now (you may have edited it). Use --force to overwrite.")
    wb = Workbook()
    w = wb.active
    w.title = "Context"
    w.append(["Key", "Value"])
    w.append(["disease", DISEASE])
    style_sheet(w, [20, 150])

    w = wb.create_sheet("Category settings")
    w.append(HEADER_SETTINGS)
    for r in SETTINGS:
        w.append(list(r))
    style_sheet(w, [22, 20, 22, 150])

    w = wb.create_sheet("Entry field changes")
    w.append(HEADER_CHANGES)
    for r in FIELD_CHANGES:
        w.append(list(r))
    style_sheet(w, [22, 18, 22, 80, 26, 20, 80])

    w = wb.create_sheet("Custom call groups")
    w.append(CALL_GROUP_HEADER)
    for r in CUSTOM_GROUPS:
        w.append(r)
    style_sheet(w, [12, 62, 60, 40, 40, 44, 40])

    w = wb.create_sheet("Custom categories")
    w.append(CAT_HEADER)
    for r in CUSTOM_CATS:
        w.append(r)
    style_sheet(w, [22, 24, 12, 12, 34, 44, 40, 62, 36, 40, 24, 40, 16, 12])

    w = wb.create_sheet("Custom entry fields")
    w.append(FIELD_HEADER)
    for r in CUSTOM_FIELDS:
        w.append(list(r))
    style_sheet(w, [22, 18, 18, 80, 60])

    w = wb.create_sheet("Rule additions")
    w.append(["Rule", "Text appended to the global rules"])
    for r in RULES:
        w.append(list(r))
    style_sheet(w, [26, 150])

    w = wb.create_sheet("Display groups")
    w.append(["Category", "Group label shown on search chips", "Match (values, or ontology ids such as HP:0002633)"])
    for r in DISPLAY_GROUPS:
        w.append(list(r))
    style_sheet(w, [22, 40, 100])

    w = wb.create_sheet("Genes")
    w.append(["Symbol", "Aliases (separate with ;)", "Notes"])
    for r in GENES:
        w.append(list(r))
    style_sheet(w, [16, 50, 100])

    write_how_to(wb, [
        "ORGANIZATION OVERLAY: VCP (DRAFT for your review)",
        "",
        "First draft written from the onboarding interview (orgs/vcp/interview.md) and a read of the 14 sample papers. Nothing here has",
        "been run through the AI yet. Edit anything: wording of the rules, the gene focus options, the three custom categories, the display",
        "groups and the gene list. Save the file, then tell Claude; the checks are free.",
        "",
        "Context: the disease description is inserted into the core instructions.",
        "Category settings / Entry field changes: how this organization differs from the shared core (see the DADA2 overlay for examples).",
        "Custom categories + Custom entry fields: the three extra categories you asked for (outcome measures and biomarkers, inhibitors and",
        "  compounds tested, model strains and cell lines). 'Call group' says which of the AI's reading passes handles each one;",
        "  group 4 is a new pass defined on the Custom call groups tab, because the API limits how many categories fit in one pass.",
        "Display groups: how phenotypes are grouped on search chips (labels must match standard Human Phenotype Ontology terms).",
        "Genes: gene names and aliases the system should treat as the same thing.",
    ])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
