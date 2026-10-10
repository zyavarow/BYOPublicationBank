"""Back-end overview for whoever runs the software: the standard categories, and what each organization has customized.

One Excel workbook:
  Summary                    one row per organization
  Standard categories        every standard category with its definition and standard options (the same for every organization)
  Custom options             per organization: options added to or removed from a standard category
  Custom wording             per organization: standard categories whose instructions were reworded
  Custom categories          per organization: categories added on top of the standard ones, in full
  Removed categories         per organization: standard categories switched off
  Genes and names            per organization: gene list with other names and roles
"""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from pubbank import orgs
from pubbank.overlay_edit import Overlay, join
from pubbank.template import DEFAULT_CORE, _read_complete_template, load_template

HEAD = PatternFill("solid", fgColor="2B6E5E")


def _core():
    wb = load_workbook(DEFAULT_CORE, data_only=True)
    return _read_complete_template(wb, DEFAULT_CORE)["categories"]


def _opts(c):
    f = next((f for f in c["fields"] if f["kind"] == "pick-list"), None)
    return (f["name"], list(f["values"])) if f else ("", [])


def collect():
    """All the rows for the sheets, as a dict of lists."""
    core = _core()
    core_by = {c["id"]: c for c in core}
    rows = {k: [] for k in ("standard", "options", "wording", "custom", "removed", "genes", "summary")}
    for c in core:
        field, vals = _opts(c)
        rows["standard"].append([c["id"], c["label"], c["answer_type"], "all papers" if c["applies_to"] == "All" else join(c["applies_to"]), c["counts"],
                                 c["not_counts"], c["interpret"], c["na_means"], c["nf_means"], c["found_in"], field, join(vals), "Yes" if c["filter_chip"] else "No"])
    rows["standard"].append(["year_authors", "Publication year, all authors, last author", "Automatic", "all papers",
                             "Filled in from PubMed or Crossref at no AI cost.", "", "", "", "", "Metadata", "", "", "No"])
    for o in orgs.list_orgs():
        path = orgs.ORGS_DIR / o["folder"] / "overlay.xlsx"
        if not path.exists():
            rows["summary"].append([o["name"], o["folder"], "no template yet", 0, 0, 0, 0, 0])
            continue
        t = load_template(path)
        ov = Overlay(path)
        custom_ids = ov.custom_ids()
        final = {c["id"]: c for c in t["categories"]}
        n_opts = n_word = 0
        for cid, c in final.items():
            if cid in custom_ids or cid not in core_by:
                continue
            core_field, core_vals = _opts(core_by[cid])
            field, vals = _opts(c)
            if field and set(vals) != set(core_vals):
                added, removed = [v for v in vals if v not in core_vals], [v for v in core_vals if v not in vals]
                rows["options"].append([o["name"], cid, c["label"], field, len(core_vals), join(added), join(removed), len(vals)])
                n_opts += 1
            for key in ("label", "counts", "not_counts", "interpret", "na_means", "nf_means", "found_in", "confidence_guide"):
                if (c[key] or "").strip() != (core_by[cid][key] or "").strip():
                    rows["wording"].append([o["name"], cid, c["label"], key, core_by[cid][key], c[key]])
                    n_word += 1
            if not c["include"]:
                rows["removed"].append([o["name"], cid, c["label"]])
        for cid in custom_ids:
            c = final.get(cid)
            if c:
                field, vals = _opts(c)
                rows["custom"].append([o["name"], cid, c["label"], c["answer_type"], "all papers" if c["applies_to"] == "All" else "primary research papers",
                                       c["counts"], c["not_counts"], field, join(vals), c["call_group"]])
        for g in t.get("genes", []):
            m = re.match(r"^\[([^\]]+)\]\s*(.*)$", g.get("notes", ""))
            rows["genes"].append([o["name"], g["symbol"], join(g["aliases"]), m.group(1) if m else "", m.group(2) if m else g.get("notes", "")])
        rows["summary"].append([o["name"], o["folder"], "template", len(core), n_opts, n_word, len(custom_ids), len([r for r in rows["removed"] if r[0] == o["name"]])])
    return rows


def workbook():
    r = collect()
    wb = Workbook()
    sheets = [
        ("Summary", ["Organization", "Folder", "Status", "Standard categories", "Standard categories with custom options", "Standard categories with custom wording",
                     "Custom categories", "Standard categories removed"], r["summary"], [34, 18, 16, 14, 18, 18, 14, 14]),
        ("Standard categories", ["ID", "Category", "Answer type", "Applies to", "What counts", "What does NOT count", "How to interpret it", "'Not applicable' means",
                                 "'Not found' means", "Usually found in", "Options field", "Standard options", "Search filter"], r["standard"],
         [20, 26, 12, 30, 50, 40, 60, 34, 36, 22, 14, 70, 10]),
        ("Custom options", ["Organization", "Category ID", "Category", "Options field", "Standard options (count)", "Options added", "Options removed", "Options now (count)"],
         r["options"], [30, 20, 26, 14, 12, 70, 40, 12]),
        ("Custom wording", ["Organization", "Category ID", "Category", "Part reworded", "Standard text", "This organization's text"], r["wording"], [30, 20, 26, 16, 60, 60]),
        ("Custom categories", ["Organization", "ID", "Category", "Answer type", "Applies to", "What counts", "What does NOT count", "Options field", "Options", "Reading pass"],
         r["custom"], [30, 24, 30, 12, 20, 60, 40, 12, 60, 10]),
        ("Removed categories", ["Organization", "Category ID", "Category"], r["removed"], [30, 24, 30]),
        ("Genes and names", ["Organization", "Gene or protein", "Other names", "Role", "Note"], r["genes"], [30, 16, 50, 24, 60]),
    ]
    for i, (name, head, data, widths) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in data:
            ws.append(row)
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
